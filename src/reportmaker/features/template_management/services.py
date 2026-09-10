"""Template management: import/export Excel templates ↔ YAML configs."""

from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from reportmaker.shared.config import load_yaml_config, save_yaml_config
from reportmaker.shared.file_utils import ensure_directory, safe_filename
from reportmaker.shared.logging import get_logger
from reportmaker.features.report_generation.models import (
    ChartConfig,
    ChartKind,
    DataSourceConfig,
    OutputFormat,
    ParameterSchema,
    ParameterType,
    ReportConfig,
    TransformConfig,
)
from reportmaker.features.template_management import xlsx_layout as L
from reportmaker.features.template_management.exceptions import (
    TemplateAlreadyExistsError,
    TemplateImportError,
    TemplateNotFoundError,
    TemplateValidationError,
)
from reportmaker.features.template_management.models import TemplateInfo, TemplateImportResult

logger = get_logger(__name__)


class TemplateManager:
    """Manages report templates stored as YAML, imported from Excel."""

    def __init__(self, templates_dir: str = "./templates"):
        self.templates_dir = Path(templates_dir)
        ensure_directory(self.templates_dir)

    # ---------- Public API ----------

    def list_templates(self) -> List[TemplateInfo]:
        """Return metadata for all stored templates."""
        infos: List[TemplateInfo] = []
        for path in sorted(self.templates_dir.glob("*.yaml")):
            name = path.stem
            description = None
            try:
                data = load_yaml_config(path)
                description = data.get("description")
            except Exception:
                pass
            infos.append(TemplateInfo(name=name, path=path, description=description))
        return infos

    def validate_template(self, name: str) -> ReportConfig:
        """Load and validate a stored template by name."""
        path = self._template_path(name)
        if not path.exists():
            raise TemplateNotFoundError(f"Template '{name}' not found at {path}")
        try:
            data = load_yaml_config(path)
            return ReportConfig(**data)
        except Exception as e:
            raise TemplateValidationError(f"Template '{name}' is invalid: {e}") from e

    def delete_template(self, name: str) -> None:
        """Delete a stored template."""
        path = self._template_path(name)
        if not path.exists():
            raise TemplateNotFoundError(f"Template '{name}' not found at {path}")
        path.unlink()
        logger.info(f"Deleted template '{name}'")

    def import_from_xlsx(
        self,
        xlsx_path: str,
        name_override: Optional[str] = None,
        force: bool = False,
    ) -> TemplateImportResult:
        """Read an Excel template, validate it, and save as YAML."""
        xlsx_path = Path(xlsx_path)
        if not xlsx_path.exists():
            return TemplateImportResult(success=False, errors=[f"File not found: {xlsx_path}"])

        try:
            config = self._xlsx_to_config(xlsx_path)
        except TemplateImportError as e:
            return TemplateImportResult(success=False, errors=[str(e)])

        # Determine template name
        template_name = name_override or config.name
        if not template_name:
            return TemplateImportResult(success=False, errors=["Template name is required"])

        safe_name = safe_filename(template_name).replace(" ", "_").lower()
        yaml_path = self._template_path(safe_name)

        if yaml_path.exists() and not force:
            return TemplateImportResult(
                success=False,
                errors=[
                    f"Template '{safe_name}' already exists at {yaml_path}. "
                    f"Use --force to overwrite."
                ],
            )

        # Save YAML
        save_yaml_config(config.model_dump(exclude_none=True), yaml_path)
        logger.info(f"Imported template '{safe_name}' -> {yaml_path}")

        return TemplateImportResult(
            success=True,
            template_name=safe_name,
            yaml_path=yaml_path,
        )

    def export_to_xlsx(self, name: str, output_path: str) -> Path:
        """Export a stored YAML template to an Excel spreadsheet."""
        path = self._template_path(name)
        if not path.exists():
            raise TemplateNotFoundError(f"Template '{name}' not found at {path}")

        data = load_yaml_config(path)
        config = ReportConfig(**data)
        return self._config_to_xlsx(config, output_path)

    def create_blank_xlsx(self, output_path: str) -> Path:
        """Create a blank template spreadsheet with headers and example rows."""
        wb = Workbook()

        # Remove default sheet
        wb.remove(wb.active)

        # Report sheet
        ws_report = wb.create_sheet(L.SHEET_REPORT)
        ws_report.append(["Field", "Value"])
        for row in L.EXAMPLE_REPORT_ROWS:
            ws_report.append(list(row))
        ws_report.column_dimensions["A"].width = 25
        ws_report.column_dimensions["B"].width = 80

        # Chart sheet
        ws_chart = wb.create_sheet(L.SHEET_CHART)
        ws_chart.append(["Field", "Value"])
        for row in L.EXAMPLE_CHART_ROWS:
            ws_chart.append(list(row))
        ws_chart.column_dimensions["A"].width = 25
        ws_chart.column_dimensions["B"].width = 40

        # Transform sheet
        ws_transform = wb.create_sheet(L.SHEET_TRANSFORM)
        ws_transform.append(["Field", "Value"])
        for row in L.EXAMPLE_TRANSFORM_ROWS:
            ws_transform.append(list(row))
        ws_transform.column_dimensions["A"].width = 25
        ws_transform.column_dimensions["B"].width = 40

        # Aggregations sheet
        ws_agg = wb.create_sheet(L.SHEET_AGGREGATIONS)
        ws_agg.append([L.AGG_HEADER_COLUMN, L.AGG_HEADER_FUNCTION])
        for row in L.EXAMPLE_AGGREGATION_ROWS:
            ws_agg.append(list(row))
        ws_agg.column_dimensions["A"].width = 30
        ws_agg.column_dimensions["B"].width = 15

        # Parameters sheet
        ws_params = wb.create_sheet(L.SHEET_PARAMETERS)
        ws_params.append([
            L.PARAM_HEADER_NAME,
            L.PARAM_HEADER_TYPE,
            L.PARAM_HEADER_REQUIRED,
            L.PARAM_HEADER_DEFAULT,
            L.PARAM_HEADER_DESCRIPTION,
        ])
        for row in L.EXAMPLE_PARAMETER_ROWS:
            ws_params.append(list(row))
        ws_params.column_dimensions["A"].width = 20
        ws_params.column_dimensions["B"].width = 12
        ws_params.column_dimensions["C"].width = 12
        ws_params.column_dimensions["D"].width = 20
        ws_params.column_dimensions["E"].width = 50

        output_path = Path(output_path)
        ensure_directory(output_path.parent)
        wb.save(output_path)
        logger.info(f"Created blank template at {output_path}")
        return output_path

    # ---------- Internal helpers ----------

    def _template_path(self, name: str) -> Path:
        safe_name = safe_filename(name).replace(" ", "_").lower()
        return self.templates_dir / f"{safe_name}.yaml"

    def _xlsx_to_config(self, xlsx_path: Path) -> ReportConfig:
        """Read an Excel file and build a ReportConfig."""
        try:
            wb = load_workbook(xlsx_path, data_only=True)
        except Exception as e:
            raise TemplateImportError(f"Could not open spreadsheet: {e}") from e

        try:
            report_kv = self._read_kv_sheet(wb[L.SHEET_REPORT]) if L.SHEET_REPORT in wb.sheetnames else {}
            chart_kv = self._read_kv_sheet(wb[L.SHEET_CHART]) if L.SHEET_CHART in wb.sheetnames else {}
            transform_kv = self._read_kv_sheet(wb[L.SHEET_TRANSFORM]) if L.SHEET_TRANSFORM in wb.sheetnames else {}
            aggregations = self._read_aggregations(wb) if L.SHEET_AGGREGATIONS in wb.sheetnames else {}
            parameters = self._read_parameters(wb) if L.SHEET_PARAMETERS in wb.sheetnames else []
        except Exception as e:
            raise TemplateImportError(f"Could not parse spreadsheet: {e}") from e

        # --- Build the config ---
        name = self._get(report_kv, L.FIELD_REPORT_NAME)
        if not name:
            raise TemplateImportError(f"'{L.FIELD_REPORT_NAME}' is required")

        db_url = self._get(report_kv, L.FIELD_DATABASE_URL)
        if not db_url:
            raise TemplateImportError(f"'{L.FIELD_DATABASE_URL}' is required")

        query = self._get(report_kv, L.FIELD_SQL_QUERY)
        if not query:
            raise TemplateImportError(f"'{L.FIELD_SQL_QUERY}' is required")

        output_format_str = self._get(report_kv, L.FIELD_OUTPUT_FORMAT, default="pdf").lower()
        try:
            output_format = OutputFormat(output_format_str)
        except ValueError:
            raise TemplateImportError(
                f"Invalid Output Format '{output_format_str}'. "
                f"Choose from: {', '.join(f.value for f in OutputFormat)}"
            )

        data_source = DataSourceConfig(url=db_url, query=query)

        # --- Transform ---
        transform: Optional[TransformConfig] = None
        group_by = self._get(transform_kv, L.FIELD_TRANSFORM_GROUP_BY)
        pivot_index = self._get(transform_kv, L.FIELD_TRANSFORM_PIVOT_INDEX)
        pivot_columns = self._get(transform_kv, L.FIELD_TRANSFORM_PIVOT_COLUMNS)
        pivot_values = self._get(transform_kv, L.FIELD_TRANSFORM_PIVOT_VALUES)

        if group_by or aggregations or (pivot_index and pivot_columns and pivot_values):
            transform = TransformConfig(
                group_by=[c.strip() for c in group_by.split(",")] if group_by else None,
                aggregations=aggregations or None,
                pivot={
                    "index": pivot_index,
                    "columns": pivot_columns,
                    "values": pivot_values,
                } if (pivot_index and pivot_columns and pivot_values) else None,
            )

        # --- Chart ---
        chart: Optional[ChartConfig] = None
        chart_kind = self._get(chart_kv, L.FIELD_CHART_KIND)
        if chart_kind:
            try:
                kind = ChartKind(chart_kind.lower())
            except ValueError:
                raise TemplateImportError(
                    f"Invalid Chart Kind '{chart_kind}'. Choose from: {', '.join(k.value for k in ChartKind)}"
                )
            x_col = self._get(chart_kv, L.FIELD_CHART_X_COLUMN)
            y_col = self._get(chart_kv, L.FIELD_CHART_Y_COLUMN)
            if not x_col or not y_col:
                raise TemplateImportError(
                    "Chart requires both 'X Column' and 'Y Column' to be filled in."
                )
            chart = ChartConfig(
                kind=kind,
                x_column=x_col,
                y_column=y_col,
                title=self._get(chart_kv, L.FIELD_CHART_TITLE),
                x_label=self._get(chart_kv, L.FIELD_CHART_X_LABEL),
                y_label=self._get(chart_kv, L.FIELD_CHART_Y_LABEL),
            )

        # --- Parameters ---
        param_schemas: Optional[List[ParameterSchema]] = None
        if parameters:
            try:
                param_schemas = [ParameterSchema(**p) for p in parameters]
            except Exception as e:
                raise TemplateImportError(f"Invalid parameters: {e}") from e

        schedule = self._get(report_kv, L.FIELD_SCHEDULE)

        try:
            return ReportConfig(
                name=name,
                description=self._get(report_kv, L.FIELD_DESCRIPTION),
                parameters=param_schemas,
                data_source=data_source,
                transform=transform,
                output=output_format,
                chart=chart,
                schedule=schedule or None,
            )
        except Exception as e:
            raise TemplateImportError(f"Configuration is invalid: {e}") from e

    def _config_to_xlsx(self, config: ReportConfig, output_path: str) -> Path:
        """Write a ReportConfig to an Excel file."""
        wb = Workbook()
        wb.remove(wb.active)

        # Report
        ws_report = wb.create_sheet(L.SHEET_REPORT)
        ws_report.append(["Field", "Value"])
        ws_report.append([L.FIELD_REPORT_NAME, config.name])
        ws_report.append([L.FIELD_DESCRIPTION, config.description or ""])
        ws_report.append([L.FIELD_DATABASE_URL, config.data_source.url])
        ws_report.append([L.FIELD_SQL_QUERY, config.data_source.query])
        ws_report.append([L.FIELD_OUTPUT_FORMAT, config.output.value])
        ws_report.append([L.FIELD_SCHEDULE, config.schedule or ""])
        ws_report.column_dimensions["A"].width = 25
        ws_report.column_dimensions["B"].width = 80

        # Chart
        ws_chart = wb.create_sheet(L.SHEET_CHART)
        ws_chart.append(["Field", "Value"])
        if config.chart:
            ws_chart.append([L.FIELD_CHART_KIND, config.chart.kind.value])
            ws_chart.append([L.FIELD_CHART_X_COLUMN, config.chart.x_column])
            ws_chart.append([L.FIELD_CHART_Y_COLUMN, config.chart.y_column])
            ws_chart.append([L.FIELD_CHART_TITLE, config.chart.title or ""])
            ws_chart.append([L.FIELD_CHART_X_LABEL, config.chart.x_label or ""])
            ws_chart.append([L.FIELD_CHART_Y_LABEL, config.chart.y_label or ""])
        else:
            ws_chart.append([L.FIELD_CHART_KIND, ""])
            ws_chart.append([L.FIELD_CHART_X_COLUMN, ""])
            ws_chart.append([L.FIELD_CHART_Y_COLUMN, ""])
            ws_chart.append([L.FIELD_CHART_TITLE, ""])
            ws_chart.append([L.FIELD_CHART_X_LABEL, ""])
            ws_chart.append([L.FIELD_CHART_Y_LABEL, ""])
        ws_chart.column_dimensions["A"].width = 25
        ws_chart.column_dimensions["B"].width = 40

        # Transform
        ws_transform = wb.create_sheet(L.SHEET_TRANSFORM)
        ws_transform.append(["Field", "Value"])
        if config.transform:
            ws_transform.append([
                L.FIELD_TRANSFORM_GROUP_BY,
                ",".join(config.transform.group_by) if config.transform.group_by else "",
            ])
            if config.transform.pivot:
                ws_transform.append([L.FIELD_TRANSFORM_PIVOT_INDEX, config.transform.pivot.get("index", "")])
                ws_transform.append([L.FIELD_TRANSFORM_PIVOT_COLUMNS, config.transform.pivot.get("columns", "")])
                ws_transform.append([L.FIELD_TRANSFORM_PIVOT_VALUES, config.transform.pivot.get("values", "")])
        else:
            ws_transform.append([L.FIELD_TRANSFORM_GROUP_BY, ""])
        ws_transform.column_dimensions["A"].width = 25
        ws_transform.column_dimensions["B"].width = 40

        # Aggregations
        ws_agg = wb.create_sheet(L.SHEET_AGGREGATIONS)
        ws_agg.append([L.AGG_HEADER_COLUMN, L.AGG_HEADER_FUNCTION])
        if config.transform and config.transform.aggregations:
            for col, func in config.transform.aggregations.items():
                ws_agg.append([col, func])
        ws_agg.column_dimensions["A"].width = 30
        ws_agg.column_dimensions["B"].width = 15

        # Parameters
        ws_params = wb.create_sheet(L.SHEET_PARAMETERS)
        ws_params.append([
            L.PARAM_HEADER_NAME,
            L.PARAM_HEADER_TYPE,
            L.PARAM_HEADER_REQUIRED,
            L.PARAM_HEADER_DEFAULT,
            L.PARAM_HEADER_DESCRIPTION,
        ])
        if config.parameters:
            for p in config.parameters:
                ws_params.append([
                    p.name,
                    p.type.value,
                    str(p.required).lower(),
                    p.default if p.default is not None else "",
                    p.description or "",
                ])
        ws_params.column_dimensions["A"].width = 20
        ws_params.column_dimensions["B"].width = 12
        ws_params.column_dimensions["C"].width = 12
        ws_params.column_dimensions["D"].width = 20
        ws_params.column_dimensions["E"].width = 50

        output_path = Path(output_path)
        ensure_directory(output_path.parent)
        wb.save(output_path)
        return output_path

    @staticmethod
    def _read_kv_sheet(ws: Worksheet) -> Dict[str, Any]:
        """Read a two-column (Field, Value) sheet into a dict."""
        result: Dict[str, Any] = {}
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return result
        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            key = str(row[0]).strip()
            value = row[1] if len(row) > 1 else None
            if value is None:
                continue
            if isinstance(value, str):
                value = value.strip()
                if not value:
                    continue
            result[key] = value
        return result

    @staticmethod
    def _read_aggregations(wb) -> Dict[str, str]:
        ws = wb[L.SHEET_AGGREGATIONS]
        result: Dict[str, str] = {}
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return result
        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            col = str(row[0]).strip()
            func = str(row[1]).strip().lower() if len(row) > 1 and row[1] else ""
            if col and func:
                result[col] = func
        return result

    @staticmethod
    def _read_parameters(wb) -> List[Dict[str, Any]]:
        ws = wb[L.SHEET_PARAMETERS]
        result: List[Dict[str, Any]] = []
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return result
        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            name = str(row[0]).strip()
            if not name:
                continue
            ptype = str(row[1]).strip().lower() if len(row) > 1 and row[1] else "string"
            required_raw = row[2] if len(row) > 2 else None
            if isinstance(required_raw, bool):
                required = required_raw
            elif isinstance(required_raw, str):
                required = required_raw.strip().lower() in ("true", "1", "yes", "y")
            else:
                required = False
            default = row[3] if len(row) > 3 else None
            if isinstance(default, str) and not default.strip():
                default = None
            description = str(row[4]).strip() if len(row) > 4 and row[4] else None
            result.append({
                "name": name,
                "type": ptype,
                "required": required,
                "default": default,
                "description": description,
            })
        return result

    @staticmethod
    def _get(kv: Dict[str, Any], key: str, default: Optional[Any] = None) -> Any:
        value = kv.get(key, default)
        if isinstance(value, str):
            value = value.strip()
            if not value:
                return default
        return value
