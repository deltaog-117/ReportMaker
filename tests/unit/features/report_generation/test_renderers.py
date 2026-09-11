"""Unit tests for the HTML and Excel renderers."""

from pathlib import Path

import pandas as pd
import pytest

from reportmaker.features.report_generation.renderers import HTMLRenderer, ExcelRenderer
from reportmaker.features.report_generation.models import (
    ChartConfig,
    ChartKind,
    DataSourceConfig,
    OutputFormat,
    ReportConfig,
)


@pytest.fixture
def sample_df():
    return pd.DataFrame({
        "product": ["Widget", "Gadget", "Doodad"],
        "revenue": [150.0, 250.0, 100.0],
    })


@pytest.fixture
def config_pdf():
    return ReportConfig(
        name="Test Report",
        description="A short description",
        data_source=DataSourceConfig(url="sqlite:///test.db", query="SELECT 1"),
        output=OutputFormat.PDF,
        chart=ChartConfig(kind=ChartKind.BAR, x_column="product", y_column="revenue"),
    )


# ---------- HTML Renderer ----------

def test_html_renderer_creates_file(tmp_path, sample_df, config_pdf):
    config = config_pdf.model_copy(update={"output": OutputFormat.HTML})
    out = tmp_path / "report.html"
    HTMLRenderer().render(sample_df, config, str(out))
    assert out.exists()
    html = out.read_text(encoding="utf-8")
    assert "Test Report" in html
    assert "Widget" in html
    assert "Gadget" in html
    # Chart is embedded as a base64 data URI
    assert "data:image/png;base64," in html


def test_html_renderer_no_data(tmp_path, config_pdf):
    config = config_pdf.model_copy(update={"output": OutputFormat.HTML, "chart": None})
    empty = pd.DataFrame({"product": [], "revenue": []})
    out = tmp_path / "empty.html"
    HTMLRenderer().render(empty, config, str(out))
    html = out.read_text(encoding="utf-8")
    assert "No data available" in html


def test_html_renderer_without_chart(tmp_path, sample_df, config_pdf):
    config = config_pdf.model_copy(update={"output": OutputFormat.HTML, "chart": None})
    out = tmp_path / "no_chart.html"
    HTMLRenderer().render(sample_df, config, str(out))
    html = out.read_text(encoding="utf-8")
    assert "Widget" in html
    assert "data:image/png;base64," not in html


# ---------- Excel Renderer ----------

def test_excel_renderer_creates_file(tmp_path, sample_df, config_pdf):
    config = config_pdf.model_copy(update={"output": OutputFormat.EXCEL})
    out = tmp_path / "report.xlsx"
    ExcelRenderer().render(sample_df, config, str(out))
    assert out.exists()

    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert "Data" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    assert "Chart" in wb.sheetnames

    ws_data = wb["Data"]
    # Header row
    assert ws_data.cell(row=1, column=1).value == "product"
    assert ws_data.cell(row=1, column=2).value == "revenue"
    # Data rows
    assert ws_data.cell(row=2, column=1).value == "Widget"
    assert ws_data.cell(row=2, column=2).value == 150.0

    ws_summary = wb["Summary"]
    # Look for the revenue total
    values = [ws_summary.cell(row=r, column=2).value for r in range(1, ws_summary.max_row + 1)]
    assert 500.0 in values  # 150 + 250 + 100


def test_excel_renderer_without_chart(tmp_path, sample_df, config_pdf):
    config = config_pdf.model_copy(update={"output": OutputFormat.EXCEL, "chart": None})
    out = tmp_path / "no_chart.xlsx"
    ExcelRenderer().render(sample_df, config, str(out))

    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert "Data" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    assert "Chart" not in wb.sheetnames


def test_excel_renderer_empty_data(tmp_path, config_pdf):
    config = config_pdf.model_copy(update={"output": OutputFormat.EXCEL})
    empty = pd.DataFrame({"product": [], "revenue": []})
    out = tmp_path / "empty.xlsx"
    ExcelRenderer().render(empty, config, str(out))
    assert out.exists()
